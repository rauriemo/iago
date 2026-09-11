"""Materialize the frozen synthetic K1 candidate corpus without touching user projects."""

import argparse
import hashlib
import json
import os
import subprocess
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


def load(source=Path("fixtures/k1/v1")):
    manifest = json.loads((source / "manifest.json").read_text(encoding="utf-8"))
    for name, digest in manifest["source_sha256"].items():
        if hashlib.sha256((source / name).read_bytes()).hexdigest() != digest:
            raise ValueError("K1 frozen source changed: " + name)
    return (
        manifest,
        json.loads((source / "documents.json").read_text(encoding="utf-8")),
        json.loads((source / "questions.json").read_text(encoding="utf-8")),
    )


def pdf(path, pages):
    output = canvas.Canvas(str(path), invariant=1)
    for text in pages:
        if text is None:
            bitmap = Image.new("RGB", (500, 150), "white")
            ImageDraw.Draw(bitmap).text((10, 30), "Synthetic scanned-only fixture", fill="black")
            output.drawImage(ImageReader(bitmap), 30, 600, width=500, height=150)
        else:
            output.drawString(40, 760, text)
        output.showPage()
    output.save()


def materialize(destination, source=Path("fixtures/k1/v1")):
    manifest, documents, questions = load(source)
    destination = Path(destination).resolve()
    destination.mkdir(parents=True, exist_ok=False)
    for record in documents:
        path = destination / record["project"] / record["path"]
        path.parent.mkdir(parents=True, exist_ok=True)
        if record["category"] == "pdf":
            pdf(path, record["content"])
        elif record["category"] == "docx":
            document = Document()
            for title, text in record["content"]:
                document.add_heading(title, level=1)
                document.add_paragraph(text)
            document.save(path)
        else:
            path.write_text("\n".join(record["content"]) + "\n", encoding="utf-8")
    root = destination / "aster"
    pdf(root / "scan-one.pdf", [None])
    pdf(root / "scan-two.pdf", [None])
    pdf(root / "mixed.pdf", ["Synthetic mixed document readable page.", None])
    (root / "malformed.docx").write_bytes(b"Synthetic malformed document")
    encrypted = PdfWriter()
    encrypted.add_blank_page(width=200, height=200)
    encrypted.encrypt("synthetic-fixture-password")
    with (root / "encrypted.pdf").open("wb") as handle:
        encrypted.write(handle)
    (root / "unsupported.bin").write_bytes(b"\x00Synthetic binary fixture")
    for name in [
        ".env.fixture",
        "credentials.txt",
        "node_modules/generated.txt",
        "dist/generated.txt",
    ]:
        path = root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("api_key = DUMMY_SYNTHETIC_K1_SECRET\n", encoding="utf-8")
    outside = destination / "outside"
    outside.mkdir()
    (outside / "private.txt").write_text("OUTSIDE_ROOT_SENTINEL", encoding="utf-8")
    (destination / "outside.txt").write_text("TRAVERSAL_SENTINEL", encoding="utf-8")
    link = root / "outside-link"
    try:
        link.symlink_to(outside, target_is_directory=True)
    except OSError:
        if os.name != "nt":
            raise
        subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link), str(outside)], capture_output=True, check=True
        )
    extras = [
        {"path": path, "status": status}
        for path, status in [
            ("scan-one.pdf", "requires_ocr"),
            ("scan-two.pdf", "requires_ocr"),
            ("mixed.pdf", "partial_requires_ocr"),
            ("malformed.docx", "unreadable"),
            ("encrypted.pdf", "encrypted"),
            ("unsupported.bin", "unsupported"),
            (".env.fixture", "excluded"),
            ("credentials.txt", "excluded"),
            ("node_modules/generated.txt", "excluded"),
            ("dist/generated.txt", "excluded"),
            ("outside-link/private.txt", "excluded"),
            ("../outside.txt", "excluded"),
        ]
    ]
    return {
        "manifest": manifest,
        "documents": documents,
        "questions": questions,
        "projects": {name: destination / name for name in manifest["projects"]},
        "extras": extras,
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    corpus = materialize(args.output)
    (args.output / "corpus.json").write_text(
        json.dumps(corpus, indent=2, default=str) + "\n", encoding="utf-8"
    )
    print(f"Created {corpus['manifest']['version']} at {args.output}; human label review pending.")
