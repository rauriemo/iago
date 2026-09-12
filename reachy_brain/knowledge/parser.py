"""Bounded document extraction in a disposable subprocess. Never executes document content."""

import io
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

TEXT_EXTENSIONS = {
    ".md",
    ".txt",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".html",
    ".css",
    ".json",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".xml",
    ".sql",
    ".rst",
    ".csv",
    ".c",
    ".cpp",
    ".h",
    ".rs",
    ".go",
    ".java",
    ".sh",
    ".ps1",
}
SECRET = re.compile(
    r"(?i)(?:api[_-]?key|access[_-]?token|secret|password|private[_-]?key)[\"']?\s*[=:]\s*[\"']?[^\s\"']{5,}|-----BEGIN .*PRIVATE KEY-----|\bsk-[A-Za-z0-9_-]{20,}"
)


def extract(path: Path) -> dict:
    if path.stat().st_size > 20 * 1024 * 1024:
        return {"status": "over_limit", "passages": []}
    suffix = path.suffix.lower()
    if suffix not in TEXT_EXTENSIONS and suffix not in {".pdf", ".docx"}:
        return {"status": "unsupported", "passages": []}
    with path.open("rb") as source:
        data = source.read(20 * 1024 * 1024 + 1)
    if len(data) > 20 * 1024 * 1024:
        return {"status": "over_limit", "passages": []}
    sections, partial = [], False
    if suffix in TEXT_EXTENSIONS:
        text = data.decode("utf-8-sig")
        if "\x00" in text:
            return {"status": "unsupported", "passages": []}
        lines = text.splitlines()
        sections = [
            ("lines", str(i + 1), "\n".join(lines[i : i + 20])) for i in range(0, len(lines), 20)
        ]
    elif suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            return {"status": "encrypted", "passages": []}
        if len(reader.pages) > 500:
            return {"status": "over_limit", "passages": []}
        for i, page in enumerate(reader.pages):
            text = page.extract_text(extraction_mode="plain") or ""
            if not text.strip():
                partial = True
            else:
                sections.append(("page", str(i + 1), text))
        if not sections:
            return {"status": "requires_ocr", "passages": []}
    elif suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if (
                sum(i.file_size for i in archive.infolist()) > 32 * 1024 * 1024
                or len(archive.infolist()) > 2000
            ):
                return {"status": "over_limit", "passages": []}
        document = Document(io.BytesIO(data))
        heading = "Document"
        for i, paragraph in enumerate(document.paragraphs):
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                heading = paragraph.text
            if paragraph.text.strip():
                sections.append(("section", f"{heading} / paragraph {i + 1}", paragraph.text))
        for i, table in enumerate(document.tables):
            sections.append(
                (
                    "section",
                    f"Table {i + 1}",
                    "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows),
                )
            )
    else:
        return {"status": "unsupported", "passages": []}
    if sum(len(s[2]) for s in sections) > 1_000_000:
        return {"status": "over_limit", "passages": []}
    if SECRET.search("\n".join(section[2] for section in sections)):
        return {"status": "excluded", "passages": []}
    passages = []
    for kind, locator, text in sections:
        for offset in range(0, len(text), 1600):
            snippet = text[offset : offset + 1600]
            if snippet.strip():
                passages.append(
                    {"locator_kind": kind, "locator": locator, "offset": offset, "text": snippet}
                )
    return {
        "status": "partial_requires_ocr" if partial else "indexed" if passages else "empty",
        "passages": passages,
    }


def main():
    try:
        result = extract(Path(sys.argv[1]))
    except Exception:
        result = {"status": "unreadable", "passages": []}
    sys.stdout.write(json.dumps(result))


if __name__ == "__main__":
    main()
