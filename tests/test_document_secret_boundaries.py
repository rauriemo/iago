"""Synthetic credential sentinels; no real credentials in fixtures or output."""

import pytest
from docx import Document

from reachy_brain.knowledge.index import ProjectIndex


@pytest.mark.features("K1", "E1")
@pytest.mark.scenario("DOCUMENT-SECRET-SECTION-BOUNDARY")
@pytest.mark.parametrize("kind", ["text_sections", "docx_paragraphs", "json_key"])
def test_secret_assignment_is_excluded_before_indexing(tmp_path, kind):
    root = tmp_path / "originals"
    root.mkdir()
    sentinel = "SYNTHETIC_BOUNDARY_CREDENTIAL"
    if kind == "docx_paragraphs":
        path = root / "brief.docx"
        document = Document()
        document.add_paragraph("password =")
        document.add_paragraph(sentinel)
        document.save(path)
    elif kind == "json_key":
        path = root / "brief.json"
        path.write_text('{"api_key": "' + sentinel + '"}')
    else:
        path = root / "brief.txt"
        path.write_text("Ordinary line\n" * 19 + "password =\n" + sentinel)
    original = path.read_bytes()
    index = ProjectIndex(tmp_path / "private" / "index.sqlite")
    project = index.add_project("Synthetic", root)
    index.activate(project)
    assert index.reindex(project)
    assert index.status(project)["coverage"] == {"excluded": 1}
    assert index.search(project, sentinel) == []
    with index.database() as db:
        assert db.execute("SELECT COUNT(*) FROM passages").fetchone()[0] == 0
    assert path.read_bytes() == original
